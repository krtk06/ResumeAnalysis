from __future__ import annotations

import io
from collections.abc import Iterator
from datetime import date

import pytest
from docx import Document
from fastapi.testclient import TestClient
from reportlab.pdfgen import canvas

from app.config import Settings
from app.main import create_app


@pytest.fixture()
def settings() -> Settings:
    return Settings(
        _env_file=None,
        environment="test",
        embedder_backend="hash",
        llm_enabled=False,
        rate_limit_per_hour=1_000,
        rate_limit_per_day=1_000,
        require_consent=True,
        log_level="WARNING",
    )


@pytest.fixture()
def client(settings: Settings) -> Iterator[TestClient]:
    app = create_app(settings)
    with TestClient(app) as test_client:
        yield test_client


def _resume_lines() -> list[str]:
    return [
        "Alex Rivera",
        "alex.rivera@example.com | +1 555 123 4567 | linkedin.com/in/alexrivera",
        "EXPERIENCE",
        "Senior Software Engineer, Acme Corp, Jan 2020 - Present",
        "Built REST APIs with Python and FastAPI, PostgreSQL, AWS, and Docker.",
        "Led a team of four engineers and cut p95 latency by 40%.",
        "Software Engineer, Beta Labs, Jun 2017 - Dec 2019",
        "Shipped a React dashboard and automated tests with PyTest.",
        "EDUCATION",
        "B.Tech in Computer Science, IIT Delhi",
        "SKILLS",
        "Python, JavaScript, TypeScript, React, FastAPI, PostgreSQL, AWS, Docker, Git, Kubernetes",
    ]


@pytest.fixture()
def resume_text() -> str:
    return "\n".join(_resume_lines())


@pytest.fixture()
def resume_lines() -> tuple[str, ...]:
    return tuple(_resume_lines())


@pytest.fixture()
def today() -> date:
    return date(2026, 10, 1)


@pytest.fixture()
def docx_bytes() -> bytes:
    document = Document()
    document.add_heading("Alex Rivera", 0)
    document.add_paragraph("alex.rivera@example.com | +1 555 123 4567 | linkedin.com/in/alexrivera")
    document.add_heading("Experience", level=1)
    document.add_paragraph("Senior Software Engineer, Acme Corp, Jan 2020 - Present")
    document.add_paragraph("Built REST APIs with Python and FastAPI, PostgreSQL, AWS, and Docker.")
    document.add_paragraph("Led a team of four engineers and cut p95 latency by 40%.")
    document.add_paragraph("Software Engineer, Beta Labs, Jun 2017 - Dec 2019")
    document.add_paragraph("Shipped a React dashboard and automated tests with PyTest.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Education"
    table.cell(0, 1).text = "B.Tech in Computer Science, IIT Delhi"
    table.cell(1, 0).text = "Skills"
    table.cell(1, 1).text = "Python, JavaScript, React, FastAPI, PostgreSQL, AWS, Docker, Git, Kubernetes"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture()
def pdf_bytes() -> bytes:
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer)
    y = 760
    for line in _resume_lines():
        pdf.drawString(60, y, line[:110])
        y -= 18
    pdf.save()
    return buffer.getvalue()


@pytest.fixture()
def blank_pdf_bytes() -> bytes:
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer)
    pdf.drawString(60, 760, " ")
    pdf.save()
    return buffer.getvalue()


@pytest.fixture()
def jd_text() -> str:
    return (
        "We are hiring a Senior Software Engineer.\n"
        "Must have Python and FastAPI experience, plus PostgreSQL.\n"
        "Required: Docker, AWS, REST API design, and Terraform.\n"
        "Nice to have: Kubernetes and React.\n"
    )
