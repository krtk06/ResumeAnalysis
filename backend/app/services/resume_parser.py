from __future__ import annotations

import io
import logging
import re
import unicodedata
import zipfile
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader

logger = logging.getLogger(__name__)


class DocumentError(ValueError):
    """Base class for safe, user-facing document errors."""


class UnsupportedDocumentError(DocumentError):
    pass


class UnreadableDocumentError(DocumentError):
    pass


class ScannedDocumentError(DocumentError):
    pass


@dataclass(frozen=True)
class ParsedDocument:
    text: str
    lines: tuple[str, ...]
    page_count: int
    source_format: str
    is_scanned: bool = False

    @property
    def char_count(self) -> int:
        return len(self.text)


_ALLOWED = {"pdf", "docx"}
_MAX_ZIP_ENTRIES = 2_000
_MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
_MAX_COMPRESSION_RATIO = 100


def _extension(filename: str | None, content_type: str | None) -> str:
    suffix = Path(filename or "").suffix.lower().lstrip(".")
    if suffix:
        return suffix
    content = (content_type or "").lower()
    if "pdf" in content:
        return "pdf"
    if "wordprocessingml" in content or "docx" in content:
        return "docx"
    return ""


def _normalise_lines(raw_lines: list[str]) -> tuple[str, ...]:
    cleaned: list[str] = []
    for raw in raw_lines:
        text = unicodedata.normalize("NFKC", raw or "")
        text = text.replace("\r\n", "\n").replace("\r", "\n")
        text = text.replace("\u00a0", " ").replace("\u2007", " ").replace("\u202f", " ")
        text = text.replace("\t", " ")
        text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)
        text = text.replace("\u200b", "").replace("\ufeff", "")
        text = re.sub(r" {2,}", " ", text).strip()
        text = re.sub(r"^[•●▪◦*\-–—]\s*", "", text).strip()
        if text:
            cleaned.append(text)
    return tuple(cleaned)


def _validate_zip(file_bytes: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as archive:
            infos = archive.infolist()
    except zipfile.BadZipFile as exc:
        raise UnsupportedDocumentError("The DOCX file is not a valid ZIP archive.") from exc

    if len(infos) > _MAX_ZIP_ENTRIES:
        raise UnreadableDocumentError("The DOCX archive contains too many entries.")

    total_uncompressed = 0
    for info in infos:
        total_uncompressed += info.file_size
        if info.file_size > _MAX_UNCOMPRESSED_BYTES:
            raise UnreadableDocumentError("The DOCX archive is too large when expanded.")
        if info.compress_size and info.file_size / max(info.compress_size, 1) > _MAX_COMPRESSION_RATIO:
            raise UnreadableDocumentError("The DOCX archive has an unsafe compression ratio.")

    if total_uncompressed > _MAX_UNCOMPRESSED_BYTES:
        raise UnreadableDocumentError("The DOCX archive is too large when expanded.")

    names = {info.filename for info in infos}
    if "[Content_Types].xml" not in names or "word/document.xml" not in names:
        raise UnsupportedDocumentError("The file is not a valid Word document.")


def _extract_pdf(file_bytes: bytes, max_pages: int) -> ParsedDocument:
    if b"%PDF-" not in file_bytes[:1024]:
        raise UnsupportedDocumentError("The file does not look like a PDF.")

    try:
        reader = PdfReader(io.BytesIO(file_bytes))
    except Exception as exc:  # pypdf raises several parser-specific errors
        raise UnreadableDocumentError("The PDF could not be read.") from exc

    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:
            raise UnreadableDocumentError("Password-protected PDFs are not supported.") from exc
        if reader.is_encrypted:
            raise UnreadableDocumentError("Password-protected PDFs are not supported.")

    try:
        page_count = len(reader.pages)
    except Exception as exc:
        raise UnreadableDocumentError("The PDF page tree could not be read.") from exc
    if page_count > max_pages:
        raise UnreadableDocumentError(f"PDFs are limited to {max_pages} pages.")

    raw_lines: list[str] = []
    for page in reader.pages:
        try:
            text = page.extract_text() or ""
        except Exception as exc:
            # One malformed page should not discard an otherwise readable PDF.
            logger.warning("Skipping unreadable PDF page: %s", exc)
            continue
        raw_lines.extend(text.splitlines())

    lines = _normalise_lines(raw_lines)
    joined = "\n".join(lines)
    return ParsedDocument(
        text=joined,
        lines=lines,
        page_count=page_count,
        source_format="pdf",
        is_scanned=len(joined) < 40 and page_count > 0,
    )


def _extract_docx(file_bytes: bytes) -> ParsedDocument:
    _validate_zip(file_bytes)

    try:
        document = Document(io.BytesIO(file_bytes))
    except Exception as exc:
        raise UnreadableDocumentError("The DOCX could not be read.") from exc

    raw_lines: list[str] = []

    # Preserve document order: paragraphs and tables can be interleaved.
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, document)
            if paragraph.text.strip():
                raw_lines.extend(paragraph.text.splitlines())
        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                line = " | ".join(cell for cell in cells if cell)
                if line:
                    raw_lines.append(line)

    for section in document.sections:
        for container in (section.header, section.footer):
            for paragraph in container.paragraphs:
                if paragraph.text.strip():
                    raw_lines.extend(paragraph.text.splitlines())

    lines = _normalise_lines(raw_lines)
    joined = "\n".join(lines)
    return ParsedDocument(
        text=joined,
        lines=lines,
        page_count=0,
        source_format="docx",
        is_scanned=False,
    )


def parse_document(
    file_bytes: bytes,
    filename: str | None,
    *,
    content_type: str | None = None,
    max_pages: int = 10,
) -> ParsedDocument:
    extension = _extension(filename, content_type)
    if extension == "doc":
        raise UnsupportedDocumentError(
            "Legacy .doc files are not supported. Save the resume as .docx or PDF and try again."
        )
    if extension not in _ALLOWED:
        raise UnsupportedDocumentError("Only PDF and DOCX files are supported.")
    if not file_bytes:
        raise UnreadableDocumentError("The uploaded file is empty.")

    if extension == "pdf":
        return _extract_pdf(file_bytes, max_pages)
    return _extract_docx(file_bytes)


def clean_text(text: str) -> str:
    """Compatibility helper that preserves line breaks."""
    return "\n".join(_normalise_lines((text or "").splitlines()))
